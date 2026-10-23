"""Lê as planilhas dos Anuários e grava a bronze em Parquet.

As tabelas do Anuário são feitas para leitura humana: cabeçalho em três linhas
com células mescladas, notas de rodapé coladas ao ano ("2024 (5)") e ao estado
("Pernambuco (4)"), e "..." onde o estado não informou. Nada disso vai para a
bronze como texto solto: cada valor sai com o ano, o crime, a medida e as notas
que o acompanham, e "..." vira nulo com o motivo ao lado.

Saídas em data/bronze/:
  estelionato_uf.parquet   um valor por edição, UF, ano, crime e medida
  notas.parquet            o texto de cada nota de rodapé, por edição
  serie_nacional.parquet   a série de estelionatos do Brasil desde 2018 (painel)
  funil.parquet            boletins, processos e presos (quadro da edição 2026)
"""

import re
from pathlib import Path

import duckdb
import openpyxl

from fontes import BRONZE, EDICOES, planilha

TITULO = "Estelionato e Estelionato por meio eletr"
NOTA = re.compile(r"\s*\((\d+)\)\s*$")

CRIMES = {"Estelionato": "total", "Estelionato por meio eletrônico": "eletronico"}
MEDIDAS = {"Ns. Absolutos": "absoluto", "Taxa": "taxa", "Taxas": "taxa", "Variação (%)": "variacao"}


def texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def separar_nota(rotulo: str) -> tuple[str, str | None]:
    """'Pernambuco (4)' -> ('Pernambuco', '4')."""
    achado = NOTA.search(rotulo)
    if not achado:
        return rotulo, None
    return rotulo[: achado.start()].strip(), achado.group(1)


def preencher(linha: list) -> list[str]:
    """Repete o rótulo de uma célula mesclada nas colunas que ela cobre."""
    saida, atual = [], ""
    for valor in linha:
        atual = texto(valor) or atual
        saida.append(atual)
    return saida


def achar_aba(livro, prefixo: str, linhas: int = 6):
    abas = []
    for aba in livro.worksheets:
        for linha in aba.iter_rows(max_row=linhas, values_only=True):
            if any(texto(v).startswith(prefixo) for v in linha):
                abas.append(aba)
                break
    return abas


def tabela_estelionato(edicao: int, aba) -> tuple[list[tuple], list[tuple]]:
    linhas = list(aba.iter_rows(values_only=True))
    # A linha dos crimes é a que começa com "Brasil e Unidades da Federação" e
    # nomeia o estelionato; as duas seguintes trazem a medida e o ano.
    i = next(
        n for n, l in enumerate(linhas)
        if texto(l[0]).startswith("Brasil e Unidades") and any(texto(v) == "Estelionato" for v in l[1:])
    )
    # Algumas edições deixam uma linha em branco entre as do cabeçalho.
    cheias = [n for n in range(i + 1, i + 6) if any(texto(v) for v in linhas[n])]
    crimes = preencher(linhas[i])
    medidas_brutas = [texto(v) for v in linhas[cheias[0]]]
    medidas = preencher(linhas[cheias[0]])
    anos = [texto(v) for v in linhas[cheias[1]]]
    i = cheias[1] - 2
    # Na variação a linha dos anos fica vazia; as demais colunas têm ano.
    colunas = []
    for c in range(1, len(crimes)):
        crime = CRIMES.get(crimes[c])
        medida = next((m for rotulo, m in MEDIDAS.items() if medidas[c].startswith(rotulo)), None)
        if crime is None or medida is None:
            continue
        # A variação ocupa uma coluna só; o preenchimento a repetiria à direita.
        if medida == "variacao" and not medidas_brutas[c]:
            continue
        ano_rotulo, ano_nota = separar_nota(anos[c]) if anos[c] else ("", None)
        ano = int(float(ano_rotulo)) if ano_rotulo else None
        colunas.append((c, crime, medida, ano, ano_nota))
    assert len(colunas) == 10, f"{edicao}: esperava 10 colunas de dado, achei {len(colunas)}"

    valores, notas = [], []
    for n in range(i + 3, len(linhas)):
        rotulo = texto(linhas[n][0])
        if not rotulo:
            continue
        if rotulo.startswith("Fonte"):
            notas.append((edicao, aba.title, "fonte", rotulo))
            for resto in linhas[n + 1:]:
                nota = texto(resto[0])
                achado = re.match(r"\((\d+|\.\.\.)\)\s*(.*)", nota, re.S)
                if achado:
                    notas.append((edicao, aba.title, achado.group(1), achado.group(2).strip()))
            break
        uf, uf_nota = separar_nota(rotulo)
        for c, crime, medida, ano, ano_nota in colunas:
            bruto = linhas[n][c]
            if isinstance(bruto, (int, float)):
                valor, situacao = float(bruto), "informado"
            elif texto(bruto) == "...":
                valor, situacao = None, "nao_informado"
            elif texto(bruto) == "-":
                valor, situacao = None, "traco"
            else:
                raise ValueError(f"{edicao} {aba.title} linha {n + 1}: valor inesperado {bruto!r}")
            valores.append((edicao, aba.title, uf, uf_nota, crime, medida, ano, ano_nota, valor, situacao, n + 1))
    return valores, notas


def serie_nacional(edicao: int, livro) -> list[tuple]:
    """A série do Brasil desde 2018 só existe no painel de estelionatos, num
    bloco de gráfico: uma linha de anos e, abaixo, a linha dos registros. A
    edição de 2023 lista o gráfico no índice, mas não traz a aba com os dados."""
    for aba in achar_aba(livro, "PAINEL"):
        linhas = list(aba.iter_rows(values_only=True))
        # Em 2024 o painel se chama "Crimes patrimoniais"; o estelionato só
        # aparece no rótulo da linha.
        if not any("stelionato" in texto(v) for l in linhas for v in l):
            continue
        for n, linha in enumerate(linhas):
            cols = [c for c, v in enumerate(linha) if isinstance(v, (int, float)) and 2018 <= v <= 2030 and float(v).is_integer()]
            if not cols or linha[cols[0]] != 2018:
                continue
            for abaixo in linhas[n + 1: n + 4]:
                numeros = [abaixo[c] for c in cols]
                rotulo = next((texto(v) for v in reversed(abaixo[: cols[0]]) if texto(v)), "")
                if all(isinstance(v, (int, float)) and v > 100_000 for v in numeros) and rotulo in ("", "Estelionato"):
                    return [(edicao, int(linha[c]), float(v)) for c, v in zip(cols, numeros)]
    print(f"{edicao}: sem série nacional na planilha")
    return []


def funil(edicao: int, livro) -> list[tuple]:
    saida = []
    for aba in achar_aba(livro, "Número de estelionatos, por Bo"):
        linhas = list(aba.iter_rows(values_only=True))
        cab = next(l for l in linhas if texto(l[0]) == "Estelionatos")
        anos = [(c, int(v)) for c, v in enumerate(cab) if isinstance(v, (int, float))]
        for l in linhas:
            if texto(l[0]) in ("Boletins de Ocorrência", "Processos Penais", "Presos"):
                saida += [(edicao, texto(l[0]), ano, float(l[c])) for c, ano in anos]
    return saida


def gravar(nome: str, colunas: str, linhas: list[tuple]) -> None:
    con = duckdb.connect()
    con.execute(f"create table t ({colunas})")
    con.executemany(f"insert into t values ({', '.join('?' * len(linhas[0]))})", linhas)
    destino = BRONZE / f"{nome}.parquet"
    con.execute(f"copy t to '{destino}' (format parquet)")
    print(f"{nome}: {len(linhas)} linhas")


def main() -> None:
    BRONZE.mkdir(parents=True, exist_ok=True)
    valores, notas, serie, etapas = [], [], [], []
    for edicao in EDICOES:
        livro = openpyxl.load_workbook(planilha(edicao), data_only=True, read_only=False)
        abas = achar_aba(livro, TITULO)
        assert len(abas) == 1, f"{edicao}: {len(abas)} tabelas de estelionato"
        v, n = tabela_estelionato(edicao, abas[0])
        valores += v
        notas += n
        serie += serie_nacional(edicao, livro)
        etapas += funil(edicao, livro)
    gravar(
        "estelionato_uf",
        "edicao int, aba varchar, uf varchar, uf_nota varchar, crime varchar, medida varchar,"
        " ano int, ano_nota varchar, valor double, situacao varchar, linha_excel int",
        valores,
    )
    gravar("notas", "edicao int, aba varchar, nota varchar, texto varchar", notas)
    gravar("serie_nacional", "edicao int, ano int, registros double", serie)
    gravar("funil", "edicao int, etapa varchar, ano int, valor double", etapas)


if __name__ == "__main__":
    main()
